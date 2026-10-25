import unittest
from io import BytesIO

from openpyxl import Workbook

from analyzer import analyze_file
from app import app


def sample_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист 1"
    sheet.append(["Отчет по обучению", None, None, None, None])
    for _ in range(6):
        sheet.append([None] * 5)
    sheet.append(["№ п/п", "Месяц", "Программа обучения", "Количество часов", "Фактически обученные"])
    sheet.append([1, "Январь", "Курс", 8, 4])
    sheet.append([2, "Январь", "Курс", 8, 3])
    sheet.append([3, "Февраль", "Курс", 8, 5])
    sheet.append([None, "Итого", None, None, 12])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class DraftTests(unittest.TestCase):
    def test_detect_and_aggregate_f01(self):
        result = analyze_file("произвольное имя.xlsx", sample_workbook())
        self.assertEqual(result["profile"].code, "F01")
        self.assertEqual(result["chart"], {"labels": ["Январь", "Февраль"], "values": [7.0, 5.0]})

    def test_f01_upload_redirects(self):
        client = app.test_client()
        response = client.post("/", data={"files": [(BytesIO(sample_workbook()), "report.xlsx")]}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, "/dashboards")

    def test_f02a_upload_redirects(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["План приема", "Количество квот", "Целевых договоров"])
        output = BytesIO()
        workbook.save(output)
        response = app.test_client().post("/", data={"files": [(BytesIO(output.getvalue()), "students.xlsx")]}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, "/dashboards")

    def test_empty_f01_template_still_redirects(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Месяц", "Программа обучения", "Количество часов", "Фактически обученные"])
        output = BytesIO()
        workbook.save(output)
        response = app.test_client().post("/", data={"files": [(BytesIO(output.getvalue()), "empty.xlsx")]}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, "/dashboards")

    def test_dashboards_placeholder(self):
        response = app.test_client().get("/dashboards")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Страница пока находится на стадии разработки".encode(), response.data)

    def test_bad_file_is_reported(self):
        client = app.test_client()
        response = client.post("/", data={"files": [(BytesIO(b"bad"), "bad.xlsx")]}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Не удалось прочитать файл".encode(), response.data)


if __name__ == "__main__":
    unittest.main()
