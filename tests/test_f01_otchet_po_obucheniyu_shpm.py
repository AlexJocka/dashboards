import io
import os
import unittest
from openpyxl import Workbook
from reports.f01_otchet_po_obucheniyu_shpm import analyze_shpm
from analyzer import analyze_file


def sample():
    w = Workbook()
    s = w.active
    s['A3'] = 'Отчет по обучению работников в рамках программы «Школа производственного мастерства»'
    s['C8'] = 'Программа обучения'
    s['E8'] = 'Фактически обученные'
    for i, month in enumerate(('июнь', 'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь')):
        row = 10 + i * 3
        s.cell(row, 2, month)
        for offset, program in enumerate(('Базовый курс', 'Продвинутый курс', 'Периодический курс')):
            s.cell(row + offset, 3, program)
    return w


class Tests(unittest.TestCase):
    def test_empty(self):
        output = analyze_shpm(sample())
        self.assertEqual(output['status'], 'empty')
        self.assertIsNone(output['kpi'][0]['value'])
        self.assertEqual(output['data_quality']['filled_numeric_cells'], 0)

    def test_values(self):
        w = sample()
        w.active['E10'] = 5
        w.active['E11'] = 2
        w.active['F10'] = 3
        output = analyze_shpm(w)
        self.assertEqual(output['kpi'][0]['value'], 7)
        self.assertEqual(output['charts'][0]['series'][0]['values'][0], 7)
        self.assertIsNone(output['charts'][0]['series'][0]['values'][1])
        stream = io.BytesIO()
        w.save(stream)
        result = analyze_file('report.xlsx', stream.getvalue(), 'F01')
        self.assertEqual(result['chart'], {'labels': ['июнь'], 'values': [7]})
        self.assertEqual(result['analytics']['report_id'], 'F01')

    def test_bad_template(self):
        w = sample()
        w.active['C10'] = 'Другая программа'
        with self.assertRaises(ValueError):
            analyze_shpm(w)

    def test_user_xlsx_when_available(self):
        filename = os.environ.get('SHPM_TEST_XLSX')
        if not filename:
            self.skipTest('SHPM_TEST_XLSX не указан')
        with open(filename, 'rb') as f:
            result = analyze_file(os.path.basename(filename), f.read(), 'F01')
        self.assertEqual(result['analytics']['report_id'], 'F01')

if __name__ == '__main__':
    unittest.main()
