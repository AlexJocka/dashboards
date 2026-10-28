"""Чтение Excel-книг, определение формы и вызов обработчиков ETL."""

from io import BytesIO

from openpyxl import load_workbook

from detector import PROFILE_BY_CODE, detect, normalize
from ETL import student_report, training_report


def analyze_file(filename: str, content: bytes, selected_code: str = "") -> dict:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Нужен файл .xlsx")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        # Осматриваем ограниченную область; личные записи не выводим на страницу.
        header_cells = []
        previews = []
        for sheet in workbook.worksheets:
            if sheet.max_row is None or sheet.max_column is None:
                sheet.calculate_dimension(force=True)
            rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), max_col=min(sheet.max_column, 35), values_only=True))
            header_cells.extend(value for row in rows for value in row)
            previews.append({
                "name": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
            })
        detected, score, candidates = detect(header_cells, workbook.sheetnames)
        profile = PROFILE_BY_CODE.get(selected_code) if selected_code else detected
        if selected_code and profile is None:
            raise ValueError("Неизвестный тип формы")
        result = {
            "filename": filename,
            "profile": profile,
            "automatic": not selected_code and detected is not None,
            "confidence": score,
            "candidates": [(code, PROFILE_BY_CODE[code].title, points) for code, points in candidates if points],
            "sheets": previews,
            "chart": None,
            "analytics": None,
            "warnings": [],
            "students": None,
            "demo": any("демо" in normalize(sheet.cell(1, 1).value).split() for sheet in workbook.worksheets),
        }
        if profile and profile.code == "F01":
            result["analytics"], result["chart"], warnings = training_report(workbook)
            result["warnings"].extend(warnings)
        elif profile and profile.code == "F02A":
            result["students"], warnings = student_report(workbook)
            result["warnings"].extend(warnings)
        elif profile:
            result["warnings"].append("Для этой формы график пока не реализован; показаны сведения о файле.")
        else:
            result["warnings"].append("Тип формы не удалось определить уверенно. Выберите его вручную и загрузите файл снова.")
        return result
    finally:
        workbook.close()
