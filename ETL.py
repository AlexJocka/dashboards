"""Извлечение и преобразование Excel-отчётов F01 (ШПМ) и F02A (студенты)."""

from collections import defaultdict
from math import isfinite
from numbers import Real

from openpyxl.utils import get_column_letter

from detector import normalize

def _student_number(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, str):
        try:
            value = float(value.replace("\u00a0", "").replace(" ", "").replace(",", "."))
        except ValueError:
            return None
    if isinstance(value, Real) and isfinite(value) and value >= 0:
        return float(value)
    return None


def student_report(workbook):
    records = []
    warnings = []
    found = False
    missing = 0
    numeric_keys = ("plan", "requested", "quota", "contracts", "within_quota", "outside_quota", "studying", "employed", "terminated", "graduates")
    for sheet in workbook.worksheets:
        header = None
        for i, row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row, 30), max_col=19, values_only=True), 1):
            names = [normalize(v) for v in row]
            if "план приема" in names[9] and "целевых договоров" in names[12] and "образовательная организация" in names[4]:
                header = i
                break
        if header is None:
            continue
        found = True
        for row in sheet.iter_rows(min_row=header + 1, max_col=19, values_only=True):
            names = [normalize(v) for v in row]
            if any(v.startswith("справочно") for v in names):
                break
            # The template contains a subheader and a row of column numbers.
            if names[4] in ("", "5") or any(v.startswith(("итого", "всего")) for v in names[:5]):
                continue
            metrics = {key: _student_number(value) for key, value in zip(numeric_keys, row[9:19])}
            missing += sum(value is None and raw not in (None, "") for value, raw in zip(metrics.values(), row[9:19]))
            records.append({
                "university": str(row[4]).strip(),
                "year": str(row[2] or "Не указан").strip(),
                "program": str(row[8] or "Не указана").strip(),
                "specialty": str(row[6] or "Не указана").strip(),
                **metrics,
            })
    if not found:
        warnings.append("Не найдена шапка сводной формы F02A. Проверьте выбранный тип отчёта.")
    elif not records:
        warnings.append("В шаблоне нет заполненных строк отчёта. Справочники не включены в расчёты.")
    if missing:
        warnings.append(f"Некорректные числовые значения пропущены: {missing}.")
    return {"records": records}, warnings


MONTHS = ('июнь', 'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь')
PROGRAMS = ('Базовый курс', 'Продвинутый курс', 'Периодический курс')
METRICS = {'hours': 4, 'completed': 5, 'in_progress': 6}


def _norm(value):
    return ' '.join(str(value or '').casefold().replace('ё', 'е').split())


def _shpm_number(value, location, warnings, integer=False):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        warnings.append(f'{location}: логическое значение вместо числа')
        return None
    try:
        if isinstance(value, Real):
            result = float(value)
        else:
            result = float(str(value).replace('\u00a0', '').replace(' ', '').replace(',', '.'))
        if not isfinite(result) or result < 0 or (integer and not result.is_integer()):
            raise ValueError
        return int(result) if integer else result
    except (ValueError, TypeError, OverflowError):
        warnings.append(f'{location}: некорректное числовое значение')
        return None


def _sum_present(values):
    present = [v for v in values if v is not None]
    return sum(present) if present else None


def analyze_shpm(workbook) -> dict:
    """Обрабатывает исходную форму F01 (колонки B:F, строки 8–30).

    Важно: значения отсутствующих ячеек — null, не ноль.
    Количество часов — сумма указанных значений по строкам, а не человеко-часы.
    """
    sheet = workbook.worksheets[0]
    if 'школа производственного' not in _norm(sheet['A3'].value):
        raise ValueError('F01: не найден заголовок ШПМ в A3')
    if _norm(sheet['C8'].value) != 'программа обучения':
        raise ValueError('F01: не найден столбец «Программа обучения» в C8')
    if 'фактически обученные' not in _norm(sheet['E8'].value):
        raise ValueError('F01: не найден заголовок «Фактически обученные» в E8')

    warnings = []
    records = []
    for month_index, month in enumerate(MONTHS):
        first_row = 10 + month_index * 3
        if _norm(sheet.cell(first_row, 2).value) != month:
            raise ValueError(f'F01: в B{first_row} ожидался месяц «{month}»')
        for offset, program in enumerate(PROGRAMS):
            row_number = first_row + offset
            if _norm(sheet.cell(row_number, 3).value) != _norm(program):
                raise ValueError(f'F01: неизвестная программа в C{row_number}')
            rec = {'month': month, 'program': program}
            for metric, column in METRICS.items():
                rec[metric] = _shpm_number(sheet.cell(row_number, column).value,
                                      f'{sheet.title}!{get_column_letter(column)}{row_number}',
                                      warnings, integer=metric != 'hours')
            records.append(rec)

    def grouping(field, keys):
        return [{field: key, **{metric: _sum_present(r[metric] for r in records if r[field] == key)
                                 for metric in METRICS}} for key in keys]

    by_month = grouping('month', MONTHS)
    by_program = grouping('program', PROGRAMS)
    totals = {metric: _sum_present(r[metric] for r in records) for metric in METRICS}
    filled = sum(r[metric] is not None for r in records for metric in METRICS)
    if filled == 0:
        warnings.append('В форме нет числовых данных для построения диаграмм.')
    elif filled < len(records) * len(METRICS):
        warnings.append('Часть показателей не заполнена: суммы рассчитаны только по имеющимся значениям.')

    return {
        'report_id': 'F01', 'title': 'Обучение по программе ШПМ',
        'status': 'empty' if filled == 0 else ('partial' if warnings else 'success'),
        'kpi': [
            {'id': 'completed', 'label': 'Завершили обучение', 'value': totals['completed']},
            {'id': 'in_progress', 'label': 'Продолжают обучение', 'value': totals['in_progress']},
            {'id': 'hours', 'label': 'Сумма часов по строкам', 'value': totals['hours']},
        ],
        'charts': [
            {'id': 'monthly', 'type': 'bar', 'title': 'Обучение по месяцам',
             'labels': list(MONTHS), 'series': [
                 {'name': 'Завершили', 'values': [r['completed'] for r in by_month]},
                 {'name': 'Продолжают', 'values': [r['in_progress'] for r in by_month]}]},
            {'id': 'programs', 'type': 'bar', 'title': 'Обучение по программам',
             'labels': list(PROGRAMS), 'series': [
                 {'name': 'Завершили', 'values': [r['completed'] for r in by_program]},
                 {'name': 'Продолжают', 'values': [r['in_progress'] for r in by_program]}]},
        ],
        'tables': [{'id': 'detail', 'title': 'Месяцы и программы', 'rows': records}],
        'data_quality': {'total_numeric_cells': 63, 'filled_numeric_cells': filled,
                         'empty_numeric_cells': 63 - filled},
        'warnings': warnings,
    }


def _display(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().replace("\n", " ")
    return text[:90] + ("…" if len(text) > 90 else "")


def _f01_chart(workbook) -> tuple[list[str], list[float], list[str]]:
    warnings = []
    totals = defaultdict(float)
    found_header = False
    for sheet in workbook.worksheets:
        month = trained = None
        for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30), values_only=True), 1):
            names = [normalize(value) for value in row]
            month = next((i for i, value in enumerate(names) if value == "месяц"), None)
            trained = next((i for i, value in enumerate(names) if "фактически обученные" in value), None)
            if month is None or trained is None:
                continue
            found_header = True
            break
        if month is None or trained is None:
            continue
        for row in sheet.iter_rows(min_row=row_number + 1, values_only=True):
            if month >= len(row) or trained >= len(row):
                continue
            period = _display(row[month])
            value = row[trained]
            if not period or "итого" in normalize(period) or "всего" in normalize(period):
                continue
            if isinstance(value, Real) and not isinstance(value, bool):
                totals[period] += float(value)
    if not found_header:
        warnings.append("Не найдены столбцы «Месяц» и «Фактически обученные».")
    if found_header and not totals:
        warnings.append("В столбце обученных нет заполненных чисел для графика.")
    return list(totals), list(totals.values()), warnings



def training_report(workbook):
    """Возвращает аналитику, совместимый график и предупреждения для F01."""
    if "школа производственного" in normalize(workbook.worksheets[0]["A3"].value):
        analytics = analyze_shpm(workbook)
        monthly = next(chart for chart in analytics["charts"] if chart["id"] == "monthly")
        points = [(label, value) for label, value in
                  zip(monthly["labels"], monthly["series"][0]["values"])
                  if value is not None]
        chart = {"labels": [label for label, _ in points],
                 "values": [value for _, value in points]} if points else None
        return analytics, chart, analytics["warnings"]
    # Поддержка прежних табличных отчётов F01 со свободным расположением шапки.
    labels, values, warnings = _f01_chart(workbook)
    return None, {"labels": labels, "values": values} if labels else None, warnings
