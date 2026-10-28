"""Аналитика F01 (ШПМ). Вход: открытая openpyxl-книга, выход: JSON-совместимый dict.
Не содержит Flask/HTML и не раскрывает персональные данные.
"""
from __future__ import annotations

from math import isfinite
from numbers import Real
from openpyxl.utils import get_column_letter

MONTHS = ('июнь', 'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь')
PROGRAMS = ('Базовый курс', 'Продвинутый курс', 'Периодический курс')
METRICS = {'hours': 4, 'completed': 5, 'in_progress': 6}


def _norm(value):
    return ' '.join(str(value or '').casefold().replace('ё', 'е').split())


def _number(value, location, warnings, integer=False):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        warnings.append(f'{location}: логическое значение вместо числа')
        return None
    try:
        if isinstance(value, Real):
            result = float(value)
        else:
            result = float(str(value).replace('\u00a0', '').replace(' ', '').replace(',', '.'))
        if not isfinite(result) or result < 0 or (integer and not result.is_integer()):
            raise ValueError
        return int(result) if integer else result
    except (ValueError, TypeError, OverflowError):
        warnings.append(f'{location}: некорректное числовое значение')
        return None


def _sum_present(values):
    present = [v for v in values if v is not None]
    return sum(present) if present else None


def analyze_shpm(workbook) -> dict:
    """Обрабатывает исходную форму F01 (колонки B:F, строки 8–30).

    Важно: значения отсутствующих ячеек — null, не ноль.
    Количество часов — сумма указанных значений по строкам, а не человеко-часы.
    """
    sheet = workbook.worksheets[0]
    if 'школа производственного' not in _norm(sheet['A3'].value):
        raise ValueError('F01: не найден заголовок ШПМ в A3')
    if _norm(sheet['C8'].value) != 'программа обучения':
        raise ValueError('F01: не найден столбец «Программа обучения» в C8')
    if 'фактически обученные' not in _norm(sheet['E8'].value):
        raise ValueError('F01: не найден заголовок «Фактически обученные» в E8')

    warnings = []
    records = []
    for month_index, month in enumerate(MONTHS):
        first_row = 10 + month_index * 3
        if _norm(sheet.cell(first_row, 2).value) != month:
            raise ValueError(f'F01: в B{first_row} ожидался месяц «{month}»')
        for offset, program in enumerate(PROGRAMS):
            row_number = first_row + offset
            if _norm(sheet.cell(row_number, 3).value) != _norm(program):
                raise ValueError(f'F01: неизвестная программа в C{row_number}')
            rec = {'month': month, 'program': program}
            for metric, column in METRICS.items():
                rec[metric] = _number(sheet.cell(row_number, column).value,
                                      f'{sheet.title}!{get_column_letter(column)}{row_number}',
                                      warnings, integer=metric != 'hours')
            records.append(rec)

    def grouping(field, keys):
        return [{field: key, **{metric: _sum_present(r[metric] for r in records if r[field] == key)
                                 for metric in METRICS}} for key in keys]

    by_month = grouping('month', MONTHS)
    by_program = grouping('program', PROGRAMS)
    totals = {metric: _sum_present(r[metric] for r in records) for metric in METRICS}
    filled = sum(r[metric] is not None for r in records for metric in METRICS)
    if filled == 0:
        warnings.append('В форме нет числовых данных для построения диаграмм.')
    elif filled < len(records) * len(METRICS):
        warnings.append('Часть показателей не заполнена: суммы рассчитаны только по имеющимся значениям.')

    return {
        'report_id': 'F01', 'title': 'Обучение по программе ШПМ',
        'status': 'empty' if filled == 0 else ('partial' if warnings else 'success'),
        'kpi': [
            {'id': 'completed', 'label': 'Завершили обучение', 'value': totals['completed']},
            {'id': 'in_progress', 'label': 'Продолжают обучение', 'value': totals['in_progress']},
            {'id': 'hours', 'label': 'Сумма часов по строкам', 'value': totals['hours']},
        ],
        'charts': [
            {'id': 'monthly', 'type': 'bar', 'title': 'Обучение по месяцам',
             'labels': list(MONTHS), 'series': [
                 {'name': 'Завершили', 'values': [r['completed'] for r in by_month]},
                 {'name': 'Продолжают', 'values': [r['in_progress'] for r in by_month]}]},
            {'id': 'programs', 'type': 'bar', 'title': 'Обучение по программам',
             'labels': list(PROGRAMS), 'series': [
                 {'name': 'Завершили', 'values': [r['completed'] for r in by_program]},
                 {'name': 'Продолжают', 'values': [r['in_progress'] for r in by_program]}]},
        ],
        'tables': [{'id': 'detail', 'title': 'Месяцы и программы', 'rows': records}],
        'data_quality': {'total_numeric_cells': 63, 'filled_numeric_cells': filled,
                         'empty_numeric_cells': 63 - filled},
        'warnings': warnings,
    }
